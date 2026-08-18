from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.utils import get_column_letter

TEXT_FIELD_TYPE = "text"
DATE_FIELD_TYPE = "date"
DECIMAL_FIELD_TYPE = "decimal"
INTEGER_FIELD_TYPE = "integer"


class DownloadTemplate:
    """Builds the blank .xlsx a user downloads to see the expected upload format.

    Shared by the transaction templates and the budget template so the downloadable
    file cannot drift away from the columns the importer actually reads.
    """

    download_headers: list[str] = []
    download_filename: str | None = None
    download_sheet_title = "Sheet1"
    download_date_fields: tuple[str, ...] = ()
    download_decimal_fields: tuple[str, ...] = ()
    download_integer_fields: tuple[str, ...] = ()
    download_date_format = "yyyy-mm-dd"
    download_decimal_format = "#,##0.00"
    download_integer_format = "0"
    download_text_format = FORMAT_TEXT
    formatted_download_rows = 1000

    def get_download_headers(self) -> list[str]:
        return self.download_headers

    def get_download_filename(self) -> str:
        return self.download_filename or "template.xlsx"

    def get_download_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_download_field_type(self, header: str) -> str:
        if header in self.download_date_fields:
            return DATE_FIELD_TYPE
        if header in self.download_decimal_fields:
            return DECIMAL_FIELD_TYPE
        if header in self.download_integer_fields:
            return INTEGER_FIELD_TYPE
        return TEXT_FIELD_TYPE

    def get_download_number_format(self, header: str) -> str:
        field_type = self.get_download_field_type(header)
        if field_type == DATE_FIELD_TYPE:
            return self.download_date_format
        if field_type == DECIMAL_FIELD_TYPE:
            return self.download_decimal_format
        if field_type == INTEGER_FIELD_TYPE:
            return self.download_integer_format
        return self.download_text_format

    def build_download_workbook(self) -> Workbook:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.download_sheet_title
        worksheet.freeze_panes = "A2"
        headers = self.get_download_headers()
        worksheet.append(headers)

        header_font = Font(bold=True)
        for column_index, header in enumerate(headers, 1):
            column_letter = get_column_letter(column_index)
            number_format = self.get_download_number_format(header)
            worksheet.cell(row=1, column=column_index).font = header_font
            worksheet.column_dimensions[column_letter].width = max(12, min(len(header) + 2, 40))
            worksheet.column_dimensions[column_letter].number_format = number_format
            for row_index in range(2, self.formatted_download_rows + 2):
                worksheet.cell(row=row_index, column=column_index).number_format = number_format

        return workbook

    def get_download_content(self) -> bytes:
        output = BytesIO()
        self.build_download_workbook().save(output)
        return output.getvalue()
