from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook

from website.data_loading.budget_template import get_budget_template
from website.data_loading.cost_line_items import (
    COST_LINE_ITEM_IMPORT_HEADERS,
    load_cost_line_items_from_file,
)
from website.models import FieldLabelOverrides, Settings
from website.tests.factories import AnalysisFactory

EXPECTED_HEADERS = [
    "Grant code *",
    "Budget line code",
    "Account code *",
    "Site code",
    "Sector code",
    "Budget line description *",
    "Total cost *",
    "LOE or Unit",
    "Months or Unit",
    "Unit cost",
    "Budget Custom Field 1",
    "Budget Custom Field 2",
    "Budget Custom Field 3",
    "Budget Custom Field 4",
    "Budget Custom Field 5",
]


def _headers(template):
    worksheet = load_workbook(BytesIO(template.get_download_content())).active
    return [cell.value for cell in worksheet[1]]


@pytest.mark.django_db
class TestBudgetTemplate:
    def test_one_column_per_importer_header_in_order(self):
        template = get_budget_template()

        assert len(template.get_download_headers()) == len(COST_LINE_ITEM_IMPORT_HEADERS)
        assert [name for _, name in template.get_download_columns()] == [
            header["name"] for header in COST_LINE_ITEM_IMPORT_HEADERS
        ]

    def test_generated_workbook_headers(self):
        assert _headers(get_budget_template()) == EXPECTED_HEADERS

    def test_required_columns_are_starred(self):
        headers = get_budget_template().get_download_headers()

        starred = [header for header in headers if header.endswith(" *")]
        assert starred == [
            "Grant code *",
            "Account code *",
            "Budget line description *",
            "Total cost *",
        ]

    def test_decimal_columns_use_the_decimal_number_format(self):
        template = get_budget_template()
        worksheet = load_workbook(BytesIO(template.get_download_content())).active

        formats = {
            cell.value: worksheet.cell(row=2, column=cell.column).number_format for cell in worksheet[1]
        }
        assert formats["Total cost *"] == template.download_decimal_format
        assert formats["Unit cost"] == template.download_decimal_format
        assert formats["Grant code *"] == template.download_text_format
        assert formats["Budget Custom Field 1"] == template.download_text_format

    def test_custom_field_headers_follow_the_label_overrides(self):
        overrides = FieldLabelOverrides.get()
        overrides.ci_dummy_field_1 = "Cost Centre"
        overrides.ci_dummy_field_1_overridden = True
        overrides.save()

        headers = get_budget_template().get_download_headers()

        assert "Cost Centre" in headers
        assert "Budget Custom Field 1" not in headers
        assert "Budget Custom Field 2" in headers

    def test_standard_budget_headers_follow_enabled_overrides(self):
        overrides = FieldLabelOverrides.get()
        for name, label in {
            "account_code": "Ledger Code",
            "sector_code": "Programme Code",
            "budget_line_description": "Budget Activity",
        }.items():
            setattr(overrides, f"ci_{name}", label)
            setattr(overrides, f"ci_{name}_overridden", True)
        overrides.save()

        expected = EXPECTED_HEADERS.copy()
        expected[2] = "Ledger Code *"
        expected[4] = "Programme Code"
        expected[5] = "Budget Activity *"
        assert _headers(get_budget_template()) == expected

        overrides.ci_account_code_overridden = False
        overrides.ci_sector_code_overridden = False
        overrides.ci_budget_line_description_overridden = False
        overrides.save()
        assert _headers(get_budget_template()) == EXPECTED_HEADERS

    @pytest.mark.parametrize("override_labels", [False, True])
    def test_generated_template_round_trips_through_the_importer(self, override_labels):
        """The file a user downloads must import cleanly once filled in."""
        if override_labels:
            overrides = FieldLabelOverrides.get()
            for name in ("account_code", "sector_code", "budget_line_description"):
                setattr(overrides, f"ci_{name}", f"Custom {name}")
                setattr(overrides, f"ci_{name}_overridden", True)
            overrides.save()
        template = get_budget_template()
        workbook = template.build_download_workbook()
        worksheet = workbook.active
        worksheet.append(
            [
                "GX922",
                "BL1",
                "9012",
                "KL89",
                "SEC1",
                "Salaries",
                1200,
                None,
                None,
                600,
                "CC-100",
                "DFID",
                "PRJ-7",
                "Phase 2",
                "Restricted",
            ]
        )
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "budget_upload_template.xlsx"

        analysis = AnalysisFactory()
        succeeded, result = load_cost_line_items_from_file(analysis, buffer)

        assert succeeded, result["errors"]
        item = analysis.cost_line_items.get()
        assert item.grant_code == "GX922"
        assert item.account_code == "9012"
        assert item.sector_code == "SEC1"
        assert item.budget_line_description == "Salaries"
        assert item.dummy_field_1 == "CC-100"
        assert item.dummy_field_5 == "Restricted"


@pytest.mark.django_db
class TestBudgetTemplateDownloadView:
    def test_download_serves_the_generated_workbook(self, analysis_workflow_with_analysis, client_with_admin):
        analysis = analysis_workflow_with_analysis.analysis
        url = reverse("budget-template-download", kwargs={"pk": analysis.pk})

        response = client_with_admin.get(url)

        assert response.status_code == 200
        assert "budget_upload_template.xlsx" in response["Content-Disposition"]
        worksheet = load_workbook(BytesIO(response.content)).active
        assert [cell.value for cell in worksheet[1]] == EXPECTED_HEADERS

    def test_load_data_links_to_the_generated_template_by_default(
        self, analysis_workflow_with_analysis, client_with_admin
    ):
        analysis = analysis_workflow_with_analysis.analysis

        response = client_with_admin.get(reverse("analysis-load-data", kwargs={"pk": analysis.pk}))

        assert response.status_code == 200
        assert reverse("budget-template-download", kwargs={"pk": analysis.pk}) in response.content.decode()

    def test_uploaded_file_overrides_the_generated_template(
        self, analysis_workflow_with_analysis, client_with_admin
    ):
        analysis = analysis_workflow_with_analysis.analysis
        site_settings = Settings.objects.first()
        site_settings.budget_upload_template = SimpleUploadedFile("bespoke.xlsx", b"bespoke")
        site_settings.save()

        response = client_with_admin.get(reverse("analysis-load-data", kwargs={"pk": analysis.pk}))
        content = response.content.decode()

        assert "bespoke" in content
        assert reverse("budget-template-download", kwargs={"pk": analysis.pk}) not in content
