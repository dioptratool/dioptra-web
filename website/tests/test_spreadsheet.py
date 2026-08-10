import io
import random
from datetime import date
from decimal import Decimal

import pytest
from django.conf import settings
from openpyxl import Workbook
from openpyxl.reader.excel import load_workbook
from openpyxl.utils import get_column_letter

from website.models import AnalysisCostType, AnalysisType, CostType
from website.tests.factories import (
    AnalysisCostTypeCategoryFactory,
    AnalysisCostTypeCategoryGrantFactory,
    AnalysisFactory,
    CategoryFactory,
    CostLineItemConfigFactory,
    CostLineItemFactory,
    CostLineItemInterventionAllocationFactory,
    CountryFactory,
    InterventionFactory,
    SubcomponentCostAllocationFactory,
    SubcomponentCostAnalysisFactory,
    UserFactory,
)
from website.models.cost_type import Indirect, Support
from website.utils.documents import _write_cost_efficiency_table, _write_other_cost_model_table
from website.views.documents import full_cost_model_spreadsheet
from website.utils.documents import (
    _fill_in_subcomponent_cost_efficiency_functions,
    _write_cost_of_each_subcomponent_per_output_metric_table,
    _write_full_cost_model_table,
    _write_metadata_table,
)


@pytest.fixture
def spreadsheet_analysis(defaults):
    intervention = InterventionFactory(
        name="My Test Intervention",
        output_metrics=[
            "NumberOfTeacherDaysOfTraining",
            "NumberOfTeacherYearsOfSupport",
        ],
    )
    analysis_type = AnalysisType.objects.create(title="My Test AnalysisType")
    country = CountryFactory(name="'Merica", code="USofA")
    owner = UserFactory(name="The True Author")
    analysis = AnalysisFactory(
        title="My Test Analysis",
        description="My Test Analysis Description",
        owner=owner,
        analysis_type=analysis_type,
        country=country,
        start_date=date(2021, 1, 1),
        end_date=date(2022, 1, 1),
        grants="DF119",
        output_count_source="My Output Count Source",
        in_kind_contributions=True,
        client_time=True,
    )
    analysis.add_intervention(
        intervention,
        parameters={
            "number_of_teachers": 40,
            "number_of_days_of_training": 80,
            "number_of_years_of_support": 10,
        },
    )
    category_1 = CategoryFactory(name="My Test Category 1")
    category_2 = CategoryFactory(name="My Test Category 2")

    # set cost_type and category as confirmed
    cost_type_category_1 = AnalysisCostTypeCategoryFactory(
        analysis=analysis,
        category=category_1,
        confirmed=True,
        cost_type=CostType.objects.get(name=settings.DEFAULT_COST_TYPE),
    )
    cost_type_category_2 = AnalysisCostTypeCategoryFactory(
        analysis=analysis,
        category=category_2,
        confirmed=True,
        cost_type=CostType.objects.get(name=settings.DEFAULT_COST_TYPE),
    )
    # combine cost_type, category, and grant
    AnalysisCostTypeCategoryGrantFactory(
        cost_type_category=cost_type_category_1,
        grant="DF119",
    )
    AnalysisCostTypeCategoryGrantFactory(
        cost_type_category=cost_type_category_2,
        grant="DF119",
    )

    line_item_1 = CostLineItemFactory(
        analysis=analysis,
        budget_line_description="My Budget Line Description 1",
        country_code="USofA",
        grant_code="DF119",
        sector_code="HEAL",
        total_cost=50000.00,
    )
    line_item_2 = CostLineItemFactory(
        analysis=analysis,
        budget_line_description="My Budget Line Description 2",
        country_code="USofA",
        grant_code="DF119",
        sector_code="HEAL",
        total_cost=25000.00,
    )
    line_item_3 = CostLineItemFactory(
        analysis=analysis,
        budget_line_description="In Kind Line Item",
        total_cost=10000.00,
    )
    line_item_4 = CostLineItemFactory(
        analysis=analysis,
        budget_line_description="Client Time Line Item",
        total_cost=5000.00,
    )
    # set allocation and category for the line items
    config1 = CostLineItemConfigFactory(
        cost_line_item=line_item_1,
        category=category_1,
    )

    CostLineItemInterventionAllocationFactory(
        cli_config=config1,
        intervention_instance=analysis.interventioninstance_set.first(),
        allocation=50,
    )

    config2 = CostLineItemConfigFactory(
        cost_line_item=line_item_2,
        category=category_2,
    )
    CostLineItemInterventionAllocationFactory(
        cli_config=config2,
        intervention_instance=analysis.interventioninstance_set.first(),
        allocation=25,
    )

    config3 = CostLineItemConfigFactory(
        cost_line_item=line_item_3,
        analysis_cost_type=AnalysisCostType.IN_KIND,
    )

    CostLineItemInterventionAllocationFactory(
        cli_config=config3,
        intervention_instance=analysis.interventioninstance_set.first(),
        allocation=75,
    )
    config4 = CostLineItemConfigFactory(
        cost_line_item=line_item_4,
        analysis_cost_type=AnalysisCostType.CLIENT_TIME,
    )

    CostLineItemInterventionAllocationFactory(
        cli_config=config4,
        intervention_instance=analysis.interventioninstance_set.first(),
        allocation=75,
    )

    analysis.calculate_output_costs()
    return analysis


class TestAnalysisSpreadsheet:
    @pytest.mark.django_db
    def test_subcomponent_metric_metadata_references_full_costs(self, spreadsheet_analysis):
        """
        Sub-component costs apply the derived split to the full output cost
        (Program + Support + Indirect), matching Insights and print.
        """
        worksheet = Workbook().active
        intervention_instance = spreadsheet_analysis.interventioninstance_set.first()
        metric_metadata = {}

        _write_cost_efficiency_table(
            worksheet,
            spreadsheet_analysis,
            intervention_instance,
            metrics_all_costs_metadata=metric_metadata,
        )

        for metric in intervention_instance.intervention.output_metric_objects():
            all_costs_row = next(
                row
                for row in range(1, worksheet.max_row + 1)
                if worksheet[f"A{row}"].value
                == f"{metric.metric_name} including Program Costs, Support Costs, Indirect Costs"
            )
            assert metric_metadata[metric.metric_name] == f"B{all_costs_row}"

    @pytest.mark.django_db
    def test_full_cost_model_derives_shared_and_skipped_subcomponent_rows(self, spreadsheet_analysis):
        intervention_instance = spreadsheet_analysis.interventioninstance_set.first()
        subcomponent_analysis = SubcomponentCostAnalysisFactory(
            intervention_instance=intervention_instance,
            subcomponent_labels=["Treatment", "Outreach"],
        )
        line_item_1 = spreadsheet_analysis.cost_line_items.get(
            budget_line_description="My Budget Line Description 1"
        )
        SubcomponentCostAllocationFactory(
            subcomponent_analysis=subcomponent_analysis,
            cli_config=line_item_1.config,
            allocations={"0": "60", "1": "40"},
        )
        # A skipped Program row with leftover values falls back to the derived split.
        line_item_2 = spreadsheet_analysis.cost_line_items.get(
            budget_line_description="My Budget Line Description 2"
        )
        SubcomponentCostAllocationFactory(
            subcomponent_analysis=subcomponent_analysis,
            cli_config=line_item_2.config,
            allocations={"0": "100", "1": "0"},
            skipped=True,
        )
        # A Support row with a stale stored allocation (migrated 2.1 data) is
        # ignored in favor of the derived split.
        support_config = CostLineItemConfigFactory(
            cost_line_item=CostLineItemFactory(
                analysis=spreadsheet_analysis,
                budget_line_description="Support Line Item",
                total_cost=10000.00,
            ),
            cost_type=CostType.objects.get(type=Support.id),
        )
        CostLineItemInterventionAllocationFactory(
            cli_config=support_config,
            intervention_instance=intervention_instance,
            allocation=100,
        )
        SubcomponentCostAllocationFactory(
            subcomponent_analysis=subcomponent_analysis,
            cli_config=support_config,
            allocations={"0": "100", "1": "0"},
        )
        # An Indirect row with no stored allocation at all.
        indirect_config = CostLineItemConfigFactory(
            cost_line_item=CostLineItemFactory(
                analysis=spreadsheet_analysis,
                budget_line_description="Indirect Line Item",
                total_cost=5000.00,
            ),
            cost_type=CostType.objects.get(type=Indirect.id),
        )
        CostLineItemInterventionAllocationFactory(
            cli_config=indirect_config,
            intervention_instance=intervention_instance,
            allocation=40,
        )

        worksheet = Workbook().active
        last_row = _write_full_cost_model_table(worksheet, spreadsheet_analysis, 1, intervention_instance)

        rows_by_description = {worksheet[f"C{row}"].value: row for row in range(3, last_row)}
        assert "In Kind Line Item" not in rows_by_description
        assert "Client Time Line Item" not in rows_by_description

        # The explicitly allocated Program row keeps the entered split.
        explicit_row = rows_by_description["My Budget Line Description 1"]
        assert worksheet[f"K{explicit_row}"].value == Decimal("0.6")
        assert worksheet[f"M{explicit_row}"].value == Decimal("0.4")

        # Shared and skipped rows carry the derived Program Cost split.
        for description in ["My Budget Line Description 2", "Support Line Item", "Indirect Line Item"]:
            row = rows_by_description[description]
            assert worksheet[f"K{row}"].value == Decimal("0.6"), description
            assert worksheet[f"M{row}"].value == Decimal("0.4"), description
            assert worksheet[f"L{row}"].value == f"=J{row} * K{row}", description
            assert worksheet[f"N{row}"].value == f"=J{row} * M{row}", description

        # The workbook's per-output formulas compute:
        #   all-costs cell * SUM(label totals) / SUMIFS(item totals with values)
        # With every row carrying a split, the denominator is the full cost
        # basis and the ratio is the derived split, so the result equals the
        # Insights amount (split * full output cost).
        item_totals = []
        treatment_totals = []
        for row in range(3, last_row):
            if worksheet[f"K{row}"].value is None:
                continue
            item_total = (
                Decimal(str(worksheet[f"G{row}"].value)) * Decimal(str(worksheet[f"H{row}"].value)) / 100
            )
            item_totals.append(item_total)
            treatment_totals.append(item_total * worksheet[f"K{row}"].value)
        assert sum(item_totals) == Decimal(
            str(spreadsheet_analysis.get_cost_output_sums_all()[intervention_instance.id])
        )
        assert sum(treatment_totals) / sum(item_totals) == Decimal("0.6")

    @pytest.mark.django_db
    def test_other_cost_table_includes_in_kind_subcomponent_allocations(self, spreadsheet_analysis):
        intervention_instance = spreadsheet_analysis.interventioninstance_set.first()
        subcomponent_analysis = SubcomponentCostAnalysisFactory(
            intervention_instance=intervention_instance,
            subcomponent_labels=["Treatment", "Outreach"],
        )
        in_kind_item = spreadsheet_analysis.in_kind_contributions_cost_line_items.get(
            budget_line_description="In Kind Line Item"
        )
        SubcomponentCostAllocationFactory(
            subcomponent_analysis=subcomponent_analysis,
            cli_config=in_kind_item.config,
            allocations={"0": "40", "1": "60"},
        )
        worksheet = Workbook().active
        other_cost_rows = {
            int(AnalysisCostType.IN_KIND): [],
            int(AnalysisCostType.CLIENT_TIME): [],
        }

        last_row = _write_other_cost_model_table(
            worksheet,
            spreadsheet_analysis,
            intervention_instance,
            other_cost_rows,
            starting_row=1,
        )

        assert [worksheet.cell(row=2, column=column).value for column in range(7, 11)] == [
            "Treatment",
            "Treatment Total",
            "Outreach",
            "Outreach Total",
        ]
        in_kind_row = other_cost_rows[int(AnalysisCostType.IN_KIND)][0]
        assert worksheet[f"G{in_kind_row}"].value == Decimal("0.4")
        assert worksheet[f"H{in_kind_row}"].value == f"=E{in_kind_row} * G{in_kind_row}"
        assert worksheet[f"I{in_kind_row}"].value == Decimal("0.6")
        assert worksheet[f"J{in_kind_row}"].value == f"=E{in_kind_row} * I{in_kind_row}"
        client_time_row = other_cost_rows[int(AnalysisCostType.CLIENT_TIME)][0]
        assert all(
            worksheet.cell(row=client_time_row, column=column).value is None for column in range(7, 11)
        )
        assert last_row == 5

    @pytest.mark.django_db
    def test_per_output_subcomponent_rows_show_amount_and_percentage(self, spreadsheet_analysis):
        """
        Item 8.1: every per-output sub-component row carries the label, the
        allocation amount (col B), and the allocation percentage (col C),
        all derived from the full-cost calculation of item 7.1. Duplicate
        labels must still reference their own cost model columns.
        """
        intervention_instance = spreadsheet_analysis.interventioninstance_set.first()
        subcomponent_analysis = SubcomponentCostAnalysisFactory(
            intervention_instance=intervention_instance,
            subcomponent_labels=["Treatment", "Treatment"],
        )
        line_item_1 = spreadsheet_analysis.cost_line_items.get(
            budget_line_description="My Budget Line Description 1"
        )
        SubcomponentCostAllocationFactory(
            subcomponent_analysis=subcomponent_analysis,
            cli_config=line_item_1.config,
            allocations={"0": "60", "1": "40"},
        )

        worksheet = Workbook().active
        metric_metadata = {}
        last_cost_efficiency_row = _write_cost_efficiency_table(
            worksheet,
            spreadsheet_analysis,
            intervention_instance,
            starting_row=1,
            metrics_all_costs_metadata=metric_metadata,
        )
        first_subcomponent_row = last_cost_efficiency_row + 2
        last_subcomponent_row = _write_cost_of_each_subcomponent_per_output_metric_table(
            worksheet,
            spreadsheet_analysis,
            intervention_instance,
            first_subcomponent_row,
        )
        first_cost_model_row = last_subcomponent_row + 2
        last_cost_model_row = _write_full_cost_model_table(
            worksheet,
            spreadsheet_analysis,
            first_cost_model_row,
            intervention_instance,
        )

        _fill_in_subcomponent_cost_efficiency_functions(
            worksheet,
            metric_metadata,
            first_subcomponent_row + 1,
            last_subcomponent_row,
            first_cost_model_row + 2,
            last_cost_model_row,
        )

        first_data = first_cost_model_row + 2
        last_data = last_cost_model_row
        metrics = intervention_instance.intervention.output_metric_objects()
        # One block per output metric: a bold title followed by one row per label.
        block_rows = 1 + len(subcomponent_analysis.subcomponent_labels)
        assert last_subcomponent_row - first_subcomponent_row == block_rows * len(metrics)

        for block_index, metric in enumerate(metrics):
            title_row = first_subcomponent_row + block_index * block_rows
            assert (
                worksheet[f"A{title_row}"].value
                == f"Cost of Each Sub-component, per {metric.cost_efficiency_unit}"
            )
            for label_idx, column in enumerate(["L", "N"]):
                label_row = title_row + 1 + label_idx
                assert worksheet[f"A{label_row}"].value == "Treatment"
                expected_percentage = (
                    f"=SUM({column}{first_data}:{column}{last_data})"
                    f" / SUMIFS(J{first_data}:J{last_data},"
                    f'{column}{first_data}:{column}{last_data}, "<>")'
                )
                assert worksheet[f"C{label_row}"].value == expected_percentage
                assert worksheet[f"C{label_row}"].number_format == "0.00%"
                assert (
                    worksheet[f"B{label_row}"].value
                    == f"={metric_metadata[metric.metric_name]} * C{label_row}"
                )
                assert worksheet[f"B{label_row}"].number_format == "$#,##0.00"
            # The amount references the metric's full-cost cell.
            referenced_row = int(metric_metadata[metric.metric_name].lstrip("B"))
            assert (
                worksheet[f"A{referenced_row}"].value
                == f"{metric.metric_name} including Program Costs, Support Costs, Indirect Costs"
            )

    @pytest.mark.django_db
    def test_per_output_subcomponent_empty_state_without_subcomponents(self, spreadsheet_analysis):
        intervention_instance = spreadsheet_analysis.interventioninstance_set.first()
        worksheet = Workbook().active

        last_row = _write_cost_of_each_subcomponent_per_output_metric_table(
            worksheet,
            spreadsheet_analysis,
            intervention_instance,
            starting_row=5,
        )

        assert worksheet["A5"].value == "No sub-component analysis for this intervention"
        assert last_row == 6
        # The fill-in pass must leave the empty state untouched.
        _fill_in_subcomponent_cost_efficiency_functions(worksheet, {}, 6, 6, 10, 10)
        assert worksheet["B5"].value is None
        assert worksheet["C5"].value is None

    @pytest.mark.django_db
    def test_subcomponent_cost_headers_use_cost_efficiency_unit(self):
        intervention = InterventionFactory(
            output_metrics=[
                "NumberOfPersonYearsOfWaterAccess",
            ],
        )
        analysis = AnalysisFactory()
        intervention_instance = analysis.add_intervention(
            intervention,
            parameters={
                "number_of_people": 100,
                "number_of_years_of_water_access": 2,
            },
        )
        SubcomponentCostAnalysisFactory(
            intervention_instance=intervention_instance,
            subcomponent_labels=["Infrastructure", "Maintenance"],
        )
        wb = Workbook()
        worksheet = wb.active

        _write_cost_of_each_subcomponent_per_output_metric_table(
            ws=worksheet,
            an_analysis=analysis,
            intervention_instance=intervention_instance,
            starting_row=1,
        )

        assert worksheet["A1"].value == "Cost of Each Sub-component, per Person-Year of Water Access"

    @pytest.mark.django_db
    def test_metadata_table_keeps_cash_parameter_numeric(self):
        intervention = InterventionFactory(
            output_metrics=[
                "ValueOfCashDistributed",
            ],
        )
        analysis = AnalysisFactory(
            currency_code="USD",
            output_count_source="My Output Count Source",
        )
        intervention_instance = analysis.add_intervention(
            intervention,
            parameters={
                "value_of_cash_distributed": 10000,
            },
        )
        wb = Workbook()
        worksheet = wb.active

        _write_metadata_table(
            ws=worksheet,
            an_analysis=analysis,
            intervention_instance=intervention_instance,
            parameter_metadata={},
            analysis_url="https://example.com/analysis/1/insights/",
        )

        assert worksheet["A9"].value == "Value of Cash Distributed"
        assert worksheet["B9"].value == 10000
        assert worksheet["B9"].number_format == '"$"#,##0.00'

    @pytest.mark.django_db
    def test_full_cost_model_spreadsheet(self, spreadsheet_analysis, rf):
        rf.user = UserFactory()

        # Run method to produce spreadsheet
        response = full_cost_model_spreadsheet(rf, spreadsheet_analysis.pk)
        assert response.status_code == 200

        virtual_workbook = io.BytesIO()

        for chunk in response.streaming_content:
            virtual_workbook.write(chunk)

        virtual_workbook.seek(0)
        wb = load_workbook(filename=virtual_workbook)

        worksheet = wb.active

        assert worksheet["A9"].value == "Number of Teachers"
        assert worksheet["A10"].value == "Number of Days of Training"
        assert worksheet["A11"].value == "Number of Years of Support"
        assert worksheet["B9"].value == 40
        assert worksheet["B10"].value == 80
        assert worksheet["B11"].value == 10
        assert worksheet["B14"].value == "The True Author"  # Author
        assert (
            worksheet["B15"].value == f"{settings.BASE_URL}/analysis/{spreadsheet_analysis.pk}/insights/"
        )  # Analysis URL
        assert worksheet["B19"].value, "=FIXED(SUM(C30) / (B9 * B10) ==  2)"
        assert worksheet["B20"].value, "=FIXED(C32 / (B9 * B10) ==  2)"
        assert (
            worksheet["B21"].value
            == "=FIXED((IFERROR(C34 / (B9 * B10), 0)) + (IFERROR(SUM(E45) / (B9 * B10), 0)), 2)"
        )
        assert worksheet["B22"].value, "=FIXED(SUM(C30) / (B9 / B11) ==  2)"
        assert worksheet["B23"].value, "=FIXED(C32 / (B9 / B11) ==  2)"
        assert (
            worksheet["B24"].value
            == "=FIXED((IFERROR(C34 / (B9 / B11), 0)) + (IFERROR(SUM(E45) / (B9 / B11), 0)), 2)"
        )
        assert worksheet["B25"].value, "=FIXED(SUM(E42) ==  2)"
        column_a_values = [cell.value for cell in worksheet["A"]]
        assert "Other Costs" in column_a_values
        assert "Other HQ Costs" not in column_a_values

    @pytest.mark.django_db
    def test_full_cost_model_spreadsheet_with_subcomponents(self, spreadsheet_analysis, rf):
        """
        Item 8.1 end to end: the generated workbook opens cleanly and shows,
        for every output metric, one sub-component block with label, amount,
        and percentage; row-level splits total 100%.
        """
        intervention_instance = spreadsheet_analysis.interventioninstance_set.first()
        subcomponent_analysis = SubcomponentCostAnalysisFactory(
            intervention_instance=intervention_instance,
            subcomponent_labels=["Treatment", "Outreach"],
        )
        line_item_1 = spreadsheet_analysis.cost_line_items.get(
            budget_line_description="My Budget Line Description 1"
        )
        SubcomponentCostAllocationFactory(
            subcomponent_analysis=subcomponent_analysis,
            cli_config=line_item_1.config,
            allocations={"0": "60", "1": "40"},
        )
        # A skipped row exports the derived split (item 7.1).
        line_item_2 = spreadsheet_analysis.cost_line_items.get(
            budget_line_description="My Budget Line Description 2"
        )
        SubcomponentCostAllocationFactory(
            subcomponent_analysis=subcomponent_analysis,
            cli_config=line_item_2.config,
            allocations={},
            skipped=True,
        )
        rf.user = UserFactory()

        response = full_cost_model_spreadsheet(rf, spreadsheet_analysis.pk)
        assert response.status_code == 200

        virtual_workbook = io.BytesIO()
        for chunk in response.streaming_content:
            virtual_workbook.write(chunk)
        virtual_workbook.seek(0)
        wb = load_workbook(filename=virtual_workbook)
        worksheet = wb.active

        title_rows = [
            row
            for row in range(1, worksheet.max_row + 1)
            if str(worksheet[f"A{row}"].value or "").startswith("Cost of Each Sub-component, per ")
        ]
        assert len(title_rows) == len(intervention_instance.intervention.output_metric_objects()) == 2

        for title_row in title_rows:
            for offset, column, label in ((1, "L", "Treatment"), (2, "N", "Outreach")):
                row = title_row + offset
                assert worksheet[f"A{row}"].value == label
                assert worksheet[f"B{row}"].value.startswith("=B")
                assert worksheet[f"B{row}"].value.endswith(f"* C{row}")
                assert worksheet[f"C{row}"].value.startswith(f"=SUM({column}")

        # Row-level splits in the cost model total 100% whether explicit or derived.
        cost_model_header_row = next(
            row
            for row in range(1, worksheet.max_row + 1)
            if worksheet[f"A{row}"].value == "Cost Type" and worksheet[f"C{row}"].value == "Cost Item"
        )
        split_rows = 0
        row = cost_model_header_row + 1
        while worksheet[f"A{row}"].value:
            if worksheet[f"K{row}"].value is not None:
                assert worksheet[f"K{row}"].value + worksheet[f"M{row}"].value == 1
                split_rows += 1
            row += 1
        assert split_rows > 0


@pytest.mark.django_db
class TestSubcomponentSectionPositions:
    """
    Property-style checks for the position-based wiring between the
    per-output sub-component section and the cost model label columns
    (item 8.1). Each seed randomizes the layout — label count and text
    (duplicates, header collisions, long and unicode names), metric count,
    section starting rows and gaps, and the mix of explicit, skipped, and
    shared rows — and the assertions check invariants instead of fixed
    cells:

    * each label row's percentage formula sums the column whose cost model
      header really is that label's own "<label> Total" column;
    * each amount formula multiplies its metric's full-cost cell by the
      row's own percentage cell; and
    * the written cells numerically reproduce `full_cost_percentages()`.
    """

    # Labels chosen to defeat text-based matching: "Item" would have
    # resolved to the cost model's fixed "Item Total" header, "Item Total"
    # and "Cost Type" collide with headers outright, plus duplicates, long,
    # and unicode names.
    LABEL_POOL = [
        "Item",
        "Item Total",
        "Cost Type",
        "Treatment",
        "Treatment",
        "Éducation à distance",
        "An extremely long sub-component label " + "x" * 100,
    ]
    METRIC_POOL = {
        "NumberOfPeople": {"number_of_people": 25},
        "NumberOfChildren": {"number_of_children": 50},
        "NumberOfParticipants": {"number_of_participants": 10},
        "NumberOfWomen": {"number_of_women": 40},
    }

    def _random_split(self, rng, count):
        cuts = sorted(rng.sample(range(0, 101), k=count - 1)) if count > 1 else []
        values = []
        previous = 0
        for cut in cuts + [100]:
            values.append(cut - previous)
            previous = cut
        return {str(idx): str(value) for idx, value in enumerate(values)}

    @pytest.mark.parametrize("seed", range(5))
    def test_randomized_layouts_keep_label_column_wiring(self, defaults, seed):
        rng = random.Random(seed)

        metric_names = rng.sample(sorted(self.METRIC_POOL), rng.randint(1, 3))
        parameters = {}
        for name in metric_names:
            parameters.update(self.METRIC_POOL[name])
        analysis = AnalysisFactory(
            in_kind_contributions=rng.choice([True, False]),
            client_time=rng.choice([True, False]),
        )
        intervention_instance = analysis.add_intervention(
            InterventionFactory(output_metrics=metric_names),
            parameters=parameters,
        )
        labels = [rng.choice(self.LABEL_POOL) for _ in range(rng.randint(1, 5))]
        subcomponent_analysis = SubcomponentCostAnalysisFactory(
            intervention_instance=intervention_instance,
            subcomponent_labels=labels,
        )

        def add_row(cost_type_type=None, allocations=None, skipped=False):
            config_kwargs = {
                "cost_line_item": CostLineItemFactory(
                    analysis=analysis,
                    total_cost=Decimal(rng.randint(100, 100000)),
                ),
            }
            if cost_type_type is not None:
                config_kwargs["cost_type"] = CostType.objects.get(type=cost_type_type)
            config = CostLineItemConfigFactory(**config_kwargs)
            CostLineItemInterventionAllocationFactory(
                cli_config=config,
                intervention_instance=intervention_instance,
                allocation=Decimal(rng.randint(1, 100)),
            )
            if allocations is not None or skipped:
                SubcomponentCostAllocationFactory(
                    subcomponent_analysis=subcomponent_analysis,
                    cli_config=config,
                    allocations=allocations or {},
                    skipped=skipped,
                )

        for _ in range(rng.randint(1, 4)):
            add_row(allocations=self._random_split(rng, len(labels)))
        if rng.random() < 0.7:
            add_row(skipped=True)
        if rng.random() < 0.7:
            add_row(cost_type_type=Support.id)
        if rng.random() < 0.7:
            add_row(cost_type_type=Indirect.id)

        analysis.calculate_output_costs()

        worksheet = Workbook().active
        metric_metadata = {}
        last_cost_efficiency_row = _write_cost_efficiency_table(
            worksheet,
            analysis,
            intervention_instance,
            starting_row=rng.randint(1, 12),
            metrics_all_costs_metadata=metric_metadata,
        )
        first_subcomponent_row = last_cost_efficiency_row + 2
        last_subcomponent_row = _write_cost_of_each_subcomponent_per_output_metric_table(
            worksheet,
            analysis,
            intervention_instance,
            first_subcomponent_row,
        )
        first_cost_model_row = last_subcomponent_row + rng.randint(2, 8)
        last_cost_model_row = _write_full_cost_model_table(
            worksheet,
            analysis,
            first_cost_model_row,
            intervention_instance,
        )
        _fill_in_subcomponent_cost_efficiency_functions(
            worksheet,
            metric_metadata,
            first_subcomponent_row + 1,
            last_subcomponent_row,
            first_cost_model_row + 2,
            last_cost_model_row,
        )

        metrics = intervention_instance.intervention.output_metric_objects()
        percentages = subcomponent_analysis.full_cost_percentages()
        header_row = first_cost_model_row + 1
        data_first = first_cost_model_row + 2
        data_last = last_cost_model_row
        block_rows = 1 + len(labels)
        assert last_subcomponent_row - first_subcomponent_row == block_rows * len(metrics)
        assert len(percentages) == len(labels)

        for metric_index, metric in enumerate(metrics):
            title_row = first_subcomponent_row + metric_index * block_rows
            assert (
                worksheet[f"A{title_row}"].value
                == f"Cost of Each Sub-component, per {metric.cost_efficiency_unit}"
            )
            for label_idx, label in enumerate(labels):
                row = title_row + 1 + label_idx
                total_column = 12 + (label_idx * 2)
                column_letter = get_column_letter(total_column)
                assert worksheet[f"A{row}"].value == label
                # The formula's column really is this label's own Total column.
                assert worksheet.cell(row=header_row, column=total_column).value == f"{label} Total"
                assert worksheet[f"C{row}"].value == (
                    f"=SUM({column_letter}{data_first}:{column_letter}{data_last})"
                    f" / SUMIFS(J{data_first}:J{data_last},"
                    f'{column_letter}{data_first}:{column_letter}{data_last}, "<>")'
                )
                assert worksheet[f"B{row}"].value == f"={metric_metadata[metric.metric_name]} * C{row}"

        # The written cells numerically reproduce the derived full-cost split.
        for label_idx, percentage in enumerate(percentages):
            numerator = Decimal(0)
            denominator = Decimal(0)
            for row in range(data_first, data_last):
                split_cell = worksheet.cell(row=row, column=11 + (label_idx * 2)).value
                if split_cell is None:
                    continue
                item_total = (
                    Decimal(str(worksheet[f"G{row}"].value)) * Decimal(str(worksheet[f"H{row}"].value)) / 100
                )
                denominator += item_total
                numerator += item_total * split_cell
            assert denominator > 0
            assert abs(numerator / denominator - Decimal(percentage) / 100) < Decimal("0.001")
