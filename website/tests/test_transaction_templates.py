import sys
import types
from datetime import date
from io import BytesIO

import pytest
from django.core.exceptions import ImproperlyConfigured
from django.test import override_settings
from openpyxl import Workbook, load_workbook

from website.data_loading.transaction_templates import TransactionTemplate, TransactionTemplateError
from website.data_loading.transaction_templates.base import (
    CANONICAL_TRANSACTION_FIELDS,
    DATE_FIELD_TYPE,
    DECIMAL_FIELD_TYPE,
    INTEGER_FIELD_TYPE,
    TEXT_FIELD_TYPE,
)
from website.data_loading.transaction_templates.registry import (
    get_enabled_transaction_template_ids,
    get_enabled_transaction_templates,
    get_transaction_template,
    get_transaction_template_choices,
)
from website.data_loading.transactions import (
    normalize_uploaded_transaction_file,
    validate_uploaded_transaction_file,
)
from website.data_loading.utils import excel_file_to_array


def _install_template_module(monkeypatch, template_id, template_cls):
    module_name = f"website.data_loading.transaction_templates.{template_id}"
    module = types.ModuleType(module_name)
    module.Template = template_cls
    monkeypatch.setitem(sys.modules, module_name, module)


def _set_active_template(monkeypatch, template_id):
    monkeypatch.setattr(
        "website.data_loading.transaction_templates.registry.get_active_transaction_template_id",
        lambda: template_id,
    )


def _analysis_with_country(country_code="SL"):
    return types.SimpleNamespace(country=types.SimpleNamespace(code=country_code))


def test_enabled_transaction_template_ids_uses_active_template(monkeypatch):
    _set_active_template(monkeypatch, "save_the_children")

    assert get_enabled_transaction_template_ids() == ["save_the_children"]


def test_transaction_template_choices_include_available_templates():
    choices = dict(get_transaction_template_choices())

    assert choices["dioptra_default"] == "Dioptra default"
    assert choices["save_the_children"] == "Save the Children"
    assert choices["catholic_relief_services"] == "Catholic Relief Services"
    assert choices["accion_contra_el_hambre"] == "Accion Contra el Hambre"
    assert (
        choices["cooperative_for_assistance_and_relief_everywhere"]
        == "Cooperative for Assistance and Relief Everywhere (CARE)"
    )
    assert choices["danish_refugee_council"] == "Danish Refugee Council"
    assert choices["mercy_corps"] == "Mercy Corps"


def test_default_transaction_template_loads(monkeypatch):
    _set_active_template(monkeypatch, "dioptra_default")
    template = get_transaction_template()

    assert template.id == "dioptra_default"
    assert template.label == "Dioptra default"
    assert template.get_download_headers() == CANONICAL_TRANSACTION_FIELDS
    assert template.get_download_filename() == "dioptra_default_transaction_template.xlsx"
    assert (
        template.get_download_content_type()
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


def test_template_packages_expose_template_interface():
    template_ids = [
        "dioptra_default",
        "save_the_children",
        "catholic_relief_services",
        "accion_contra_el_hambre",
        "cooperative_for_assistance_and_relief_everywhere",
        "danish_refugee_council",
        "mercy_corps",
    ]

    for template_id in template_ids:
        module = __import__(
            f"website.data_loading.transaction_templates.{template_id}",
            fromlist=["Template"],
        )

        assert module.__all__ == ["Template"]
        assert module.Template.__module__.endswith(f"{template_id}.template")


@pytest.mark.parametrize(
    "template_id",
    [
        "dioptra_default",
        "save_the_children",
        "catholic_relief_services",
        "accion_contra_el_hambre",
        "cooperative_for_assistance_and_relief_everywhere",
        "danish_refugee_council",
        "mercy_corps",
    ],
)
def test_transaction_templates_define_all_validation_field_sources(monkeypatch, template_id):
    _set_active_template(monkeypatch, template_id)
    template = get_transaction_template()

    assert set(template.canonical_field_sources) == set(CANONICAL_TRANSACTION_FIELDS)
    assert set(template.get_validation_field_labels([template.get_download_headers()])) == set(
        CANONICAL_TRANSACTION_FIELDS
    )


@pytest.mark.parametrize(
    "template_id,date_header,decimal_header,text_header",
    [
        ("dioptra_default", "transaction_date", "amount", "grant_code"),
        ("save_the_children", "column_5", "column_6", "column_1"),
        ("catholic_relief_services", "Date", "Amount", "Grant_code"),
        ("accion_contra_el_hambre", "trans_date", "amount_eur", "contract"),
        (
            "cooperative_for_assistance_and_relief_everywhere",
            "journal_date",
            "bu_amount",
            "grant_code",
        ),
        ("danish_refugee_council", "column_6", "column_18", "column_9"),
        ("mercy_corps", "posting_date", "usd_amount", "posting_period"),
    ],
)
def test_transaction_template_download_workbook_formats_typed_fields(
    monkeypatch,
    template_id,
    date_header,
    decimal_header,
    text_header,
):
    _set_active_template(monkeypatch, template_id)
    template = get_transaction_template()

    assert template.get_download_field_type(date_header) == DATE_FIELD_TYPE
    assert template.get_download_field_type(decimal_header) == DECIMAL_FIELD_TYPE
    assert template.get_download_field_type(text_header) == TEXT_FIELD_TYPE

    workbook = load_workbook(BytesIO(template.get_download_content()))
    worksheet = workbook.active
    headers = template.get_download_headers()
    date_column = headers.index(date_header) + 1
    decimal_column = headers.index(decimal_header) + 1
    text_column = headers.index(text_header) + 1

    assert [cell.value for cell in worksheet[1]] == headers
    assert worksheet.cell(row=2, column=date_column).number_format == template.download_date_format
    assert worksheet.cell(row=2, column=decimal_column).number_format == template.download_decimal_format
    assert worksheet.cell(row=2, column=text_column).number_format == template.download_text_format


def test_downloaded_xlsx_template_upload_formats_date_cells(monkeypatch):
    _set_active_template(monkeypatch, "dioptra_default")
    template = get_transaction_template()
    workbook = template.build_download_workbook()
    worksheet = workbook.active
    values = [
        date(2015, 1, 1),
        "JO",
        "100",
        "200",
        "300",
        "",
        "",
        "",
        "",
        "USD",
        "Budget line",
        "123.45",
    ]
    for column_index, value in enumerate(values, 1):
        worksheet.cell(row=2, column=column_index).value = value
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    succeeded, normalized = normalize_uploaded_transaction_file(
        excel_file_to_array(output),
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["transaction_date"] == "2015-01-01"


def test_excel_file_to_array_streams_only_active_xlsx_worksheet(monkeypatch):
    workbook = Workbook()
    inactive_worksheet = workbook.active
    inactive_worksheet.title = "Inactive"
    inactive_worksheet.append(["inactive header"])
    inactive_worksheet.append(["inactive value"])

    active_worksheet = workbook.create_sheet("Active")
    active_worksheet.append(["active header", "amount"])
    active_worksheet.append(["active value", 123])
    workbook.active = workbook.index(active_worksheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    load_options = {}

    def tracked_load_workbook(filename, **kwargs):
        load_options.update(kwargs)
        return load_workbook(filename, **kwargs)

    monkeypatch.setattr("website.data_loading.utils.load_workbook", tracked_load_workbook)

    assert excel_file_to_array(output) == [
        ["active header", "amount"],
        ["active value", "123"],
    ]
    assert load_options == {
        "data_only": True,
        "read_only": True,
        "keep_links": False,
    }


@pytest.mark.parametrize(
    "template_id,source_key,primary_format,source_date,expected_date,row_values",
    [
        (
            "dioptra_default",
            "transaction_date",
            "%Y-%m-%d",
            "2015-01-01",
            "2015-01-01",
            {
                "country_code": "JO",
                "grant_code": "100",
                "account_code": "300",
                "currency_code": "USD",
                "budget_line_description": "Budget line",
                "amount": "123.45",
            },
        ),
        (
            "save_the_children",
            "column_5",
            "%Y%m",
            "202605",
            "2026-05-01",
            {
                "column_1": "SC-123",
                "column_3": "Medical supplies",
                "column_4": "5501",
                "column_6": "1234.56",
            },
        ),
        (
            "catholic_relief_services",
            "date",
            "%Y-%m-%d",
            "2025-11-30",
            "2025-11-30",
            {
                "Country_code": "2624",
                "Grant_code": "30950",
                "Account_code": "601101",
                "Currency_code": "USD",
                "Budget_line_description": "Allocated Direct Salaries - International",
                "Amount": "131.8",
            },
        ),
        (
            "accion_contra_el_hambre",
            "trans_date",
            "%m/%d/%Y",
            "10/07/2024",
            "2024-10-07",
            {
                "contract": "MLB2AS",
                "account": "600100",
                "amount_eur": "2458.70",
                "dim_1": "ML",
            },
        ),
        (
            "cooperative_for_assistance_and_relief_everywhere",
            "journal_date",
            "%m/%d/%Y",
            "03/31/2024",
            "2024-03-31",
            {
                "country_code": "KEN01",
                "grant_code": "EF789",
                "account": "501100",
                "currency_cd": "USD",
                "budget_line_description": "5 - LER program manager",
                "bu_amount": "103.14",
            },
        ),
        (
            "danish_refugee_council",
            "column_6",
            "%Y-%m-%d",
            "2024-09-30",
            "2024-09-30",
            {
                "column_8": "BGD",
                "column_9": "Gr00006009",
                "column_11": "23011",
                "column_16": "BDT",
                "column_17": "BGD-082554 - 1.B.1 - Head of Programme (1)",
                "column_18": "7823",
            },
        ),
        (
            "mercy_corps",
            "posting_date",
            "%Y-%m-%d",
            "2023-04-01",
            "2023-04-01",
            {
                "fund_no": "999dummy",
                "dept_name": "ESTORIA DEPARTMENT",
                "g_l_account_no": "6910",
                "lin_name": "Office Costs",
                "usd_amount": "9.85",
            },
        ),
    ],
)
def test_transaction_templates_parse_configured_source_date_formats(
    monkeypatch,
    template_id,
    source_key,
    primary_format,
    source_date,
    expected_date,
    row_values,
):
    _set_active_template(monkeypatch, template_id)
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    date_header = template.get_source_field_display_name(source_key)
    source_row.update(row_values)
    source_row[date_header] = source_date

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=_analysis_with_country(),
    )

    assert template.source_date_formats[source_key][0] == primary_format
    assert succeeded
    assert normalized[0]["transaction_date"] == expected_date


def test_save_the_children_transaction_template_keeps_day_month_year_period_support(monkeypatch):
    _set_active_template(monkeypatch, "save_the_children")
    template = get_transaction_template()

    succeeded, normalized = normalize_uploaded_transaction_file(
        [["SC-123", "", "Medical supplies", "5501", "29/05/2026", "1234.56"]],
        analysis=_analysis_with_country(),
    )

    assert "%d/%m/%Y" in template.source_date_formats["column_5"]
    assert succeeded
    assert normalized[0]["transaction_date"] == "2026-05-29"


def test_transaction_templates_parse_non_padded_month_day_source_date_formats(monkeypatch):
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "journal_date": "3/1/2024",
            "country_code": "KEN01",
            "grant_code": "EF789",
            "account": "501100",
            "currency_cd": "USD",
            "budget_line_description": "5 - LER program manager",
            "bu_amount": "103.14",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=object(),
    )

    assert "%m/%d/%Y" in template.source_date_formats["journal_date"]
    assert succeeded
    assert normalized[0]["transaction_date"] == "2024-03-01"


def test_transaction_templates_parse_iso_source_date_fallback(monkeypatch):
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    template = get_transaction_template()

    assert template.source_date_formats["journal_date"] == ("%m/%d/%Y",)
    assert template.parse_source_date("journal_date", "2024-03-31") == "2024-03-31"


def test_transaction_template_reports_source_date_format_errors(monkeypatch):
    _set_active_template(monkeypatch, "save_the_children")

    succeeded, errors = normalize_uploaded_transaction_file(
        [["SC-123", "", "Medical supplies", "5501", "05-29-2026", "1234.56"]],
        analysis=_analysis_with_country(),
    )

    assert not succeeded
    # "Row 2" for what is physically row 1: source_rows_from_upload enumerates from 2 because
    # rows_with_source_headers has prepended the synthetic header row. Same in the DRC template.
    assert errors == ["Row 2: column_5 must use date format %Y%m or %d/%m/%Y or %Y-%m-%d (got 05-29-2026)"]


def test_mercy_corps_transaction_template_keeps_period_and_quarter_as_text(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()

    assert template.get_download_field_type("posting_date") == DATE_FIELD_TYPE
    assert template.get_download_field_type("document_date") == DATE_FIELD_TYPE
    assert template.get_download_field_type("processing_date") == DATE_FIELD_TYPE
    assert template.get_download_field_type("dept_code") == INTEGER_FIELD_TYPE
    assert template.get_download_field_type("entry_no") == INTEGER_FIELD_TYPE
    assert template.get_download_field_type("posting_period") == TEXT_FIELD_TYPE
    assert template.get_download_field_type("calendar_quarter") == TEXT_FIELD_TYPE
    assert "posting_period" not in template.source_date_formats
    assert "calendar_quarter" not in template.source_date_formats

    workbook = load_workbook(BytesIO(template.get_download_content()))
    worksheet = workbook.active
    headers = template.get_download_headers()
    dept_code_column = headers.index("dept_code") + 1
    entry_no_column = headers.index("entry_no") + 1

    assert worksheet.cell(row=2, column=dept_code_column).number_format == template.download_integer_format
    assert worksheet.cell(row=2, column=entry_no_column).number_format == template.download_integer_format


def test_default_transaction_template_ignores_download_header_row(monkeypatch):
    _set_active_template(monkeypatch, "dioptra_default")
    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            CANONICAL_TRANSACTION_FIELDS,
            [
                "2015-01-01",
                "JO",
                "100.0",
                "200",
                "300",
                "",
                "",
                "",
                "",
                "USD",
                "Budget line",
                "123.45",
            ],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["transaction_date"] == "2015-01-01"
    assert normalized[0]["grant_code"] == "100"


def test_transaction_template_normalizes_amount_thousands_separators(monkeypatch):
    _set_active_template(monkeypatch, "dioptra_default")
    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            CANONICAL_TRANSACTION_FIELDS,
            [
                "2015-01-01",
                "JO",
                "100",
                "200",
                "300",
                "",
                "",
                "",
                "",
                "USD",
                "Budget line",
                "1,000.00",
            ],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["amount"] == "1000.00"


# Layouts whose source files carry no currency column at all.
CURRENCY_LESS_LAYOUTS = [
    (
        "save_the_children",
        {
            "column_1": "SC-123",
            "column_3": "Medical supplies",
            "column_4": "5501",
            "column_5": "202605",
            "column_6": "1234.56",
        },
    ),
    (
        "accion_contra_el_hambre",
        {
            "contract": "MLB2AS",
            "trans_date": "10/07/2024",
            "account": "600100",
            "amount_eur": "2458.70",
            "dim_1": "ML",
        },
    ),
    (
        "mercy_corps",
        {
            "posting_date": "2023-04-01",
            "fund_no": "999dummy",
            "g_l_account_no": "6910",
            "g_l_account_name": "Office Costs",
            "lin_name": "Office Costs",
            "usd_amount": "9.85",
        },
    ),
]


def _normalize_source_row(template, row_values):
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(row_values)
    return normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=_analysis_with_country(),
    )


@override_settings(ISO_CURRENCY_CODE="GBP")
@pytest.mark.parametrize("template_id,row_values", CURRENCY_LESS_LAYOUTS)
def test_layout_without_a_currency_column_takes_the_instance_currency(monkeypatch, template_id, row_values):
    _set_active_template(monkeypatch, template_id)
    template = get_transaction_template()

    assert template.canonical_field_sources["currency_code"] is None

    succeeded, normalized = _normalize_source_row(template, row_values)

    assert succeeded, normalized
    assert normalized[0]["currency_code"] == "GBP"


@override_settings(ISO_CURRENCY_CODE="none")
@pytest.mark.parametrize("template_id,row_values", CURRENCY_LESS_LAYOUTS)
def test_layout_without_a_currency_column_stays_blank_without_an_instance_currency(
    monkeypatch, template_id, row_values
):
    """Nothing can supply a currency here, so the rows simply carry none."""
    _set_active_template(monkeypatch, template_id)
    template = get_transaction_template()

    succeeded, normalized = _normalize_source_row(template, row_values)

    assert succeeded, normalized
    assert normalized[0]["currency_code"] == ""


@override_settings(ISO_CURRENCY_CODE="GBP")
def test_transaction_upload_falls_back_to_the_instance_currency(monkeypatch):
    """A row that carries no currency is imported in the instance currency."""
    _set_active_template(monkeypatch, "dioptra_default")
    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            CANONICAL_TRANSACTION_FIELDS,
            ["2015-01-01", "JO", "100", "200", "300", "", "", "", "", "", "Budget line", "123.45"],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["currency_code"] == "GBP"


@override_settings(ISO_CURRENCY_CODE="GBP")
def test_transaction_upload_keeps_the_currency_the_file_supplies(monkeypatch):
    _set_active_template(monkeypatch, "dioptra_default")
    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            CANONICAL_TRANSACTION_FIELDS,
            ["2015-01-01", "JO", "100", "200", "300", "", "", "", "", "BDT", "Budget line", "123.45"],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["currency_code"] == "BDT"


@override_settings(ISO_CURRENCY_CODE="none")
def test_transaction_upload_leaves_currency_blank_without_an_instance_currency(monkeypatch):
    _set_active_template(monkeypatch, "dioptra_default")
    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            CANONICAL_TRANSACTION_FIELDS,
            ["2015-01-01", "JO", "100", "200", "300", "", "", "", "", "", "Budget line", "123.45"],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["currency_code"] == ""


@override_settings(ISO_CURRENCY_CODE="EUR")
def test_currency_column_is_optional_when_the_instance_has_a_currency(monkeypatch):
    """CARE files identify columns by header, so the currency column can be left out."""
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    template = get_transaction_template()
    headers = [header for header in template.get_download_headers() if header != "currency_cd"]
    source_row = dict.fromkeys(headers, "")
    source_row.update(
        {
            "journal_date": "03/31/2024",
            "country_code": "KEN01",
            "grant_code": "EF789",
            "account": "501100",
            "budget_line_description": "5 - LER program manager",
            "bu_amount": "103.14",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [headers, [source_row[header] for header in headers]],
        analysis=object(),
    )

    assert succeeded, normalized
    assert normalized[0]["currency_code"] == "EUR"


@override_settings(ISO_CURRENCY_CODE="none")
def test_currency_column_is_required_without_an_instance_currency(monkeypatch):
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    template = get_transaction_template()
    headers = [header for header in template.get_download_headers() if header != "currency_cd"]

    succeeded, errors = normalize_uploaded_transaction_file(
        [headers, ["" for _ in headers]],
        analysis=object(),
    )

    assert not succeeded
    assert "currency_cd" in errors[0]


def test_transaction_template_must_be_active(monkeypatch):
    _set_active_template(monkeypatch, "dioptra_default")

    with pytest.raises(ImproperlyConfigured, match="not active"):
        get_transaction_template("save_the_children")


def test_transaction_template_choices_can_include_plugin_modules(monkeypatch):
    class Template(TransactionTemplate):
        id = "fake_partner"
        label = "Fake Partner"

    _install_template_module(monkeypatch, "fake_partner", Template)
    monkeypatch.setattr(
        "website.data_loading.transaction_templates.registry.get_available_transaction_template_ids",
        lambda: ["dioptra_default", "fake_partner"],
    )

    choices = get_transaction_template_choices()

    assert choices == [("dioptra_default", "Dioptra default"), ("fake_partner", "Fake Partner")]


def test_enabled_transaction_templates_returns_only_active_template(monkeypatch):
    _set_active_template(monkeypatch, "save_the_children")

    templates = get_enabled_transaction_templates()

    assert [template.id for template in templates] == ["save_the_children"]


def test_normalize_uploaded_transaction_file_uses_selected_template(monkeypatch):
    class Template(TransactionTemplate):
        id = "fake_partner"
        label = "Fake Partner"

        def normalize_rows(self, rows, analysis):
            assert rows == [{"source_field": "source value"}]
            return [
                {
                    "transaction_date": "2015-01-01",
                    "country_code": "JO",
                    "grant_code": "100.0",
                    "budget_line_code": "200",
                    "account_code": "300",
                    "currency_code": "USD",
                    "budget_line_description": "Budget line",
                    "amount": "123.45",
                }
            ]

    _install_template_module(monkeypatch, "fake_partner", Template)
    _set_active_template(monkeypatch, "fake_partner")

    succeeded, normalized = normalize_uploaded_transaction_file(
        [["Source Field"], ["source value"]],
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["grant_code"] == "100"


def test_normalize_uploaded_transaction_file_returns_template_errors(monkeypatch):
    class Template(TransactionTemplate):
        id = "error_partner"
        label = "Error Partner"

        def normalize_rows(self, rows, analysis):
            raise TransactionTemplateError(["Row 1: Missing required partner field"])

    _install_template_module(monkeypatch, "error_partner", Template)
    _set_active_template(monkeypatch, "error_partner")

    succeeded, errors = normalize_uploaded_transaction_file(
        [["not", "canonical"]],
        analysis=object(),
    )

    assert not succeeded
    assert errors == ["Row 1: Missing required partner field"]


def test_save_the_children_transaction_template_loads(monkeypatch):
    _set_active_template(monkeypatch, "save_the_children")
    template = get_transaction_template()

    assert template.id == "save_the_children"
    assert template.label == "Save the Children"
    assert template.get_download_headers() == [f"column_{n}" for n in range(1, 13)]


@override_settings(ISO_CURRENCY_CODE="USD")
def test_save_the_children_transaction_template_normalizes_positional_rows(monkeypatch):
    _set_active_template(monkeypatch, "save_the_children")

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            [
                "57801594",  # column_1  grant_code
                "4021014ON14A",  # column_2  budget_line_code
                "DM - ON14A-Other nutrition",  # column_3  budget_line_description
                "52010",  # column_4  account_code
                "202306",  # column_5  transaction_date
                "39.88",  # column_6  amount
                "NUT",  # column_7  sector_code
                "Finance coordinator",  # column_8  custom field 1
                "6-CAM Support costs -2022",  # column_9  custom field 2
                "",  # column_10 custom field 3
                "6-CAM Support costs -2022",  # column_11 custom field 4
                "",  # column_12 custom field 5
            ]
        ],
        analysis=_analysis_with_country("SL"),
    )

    assert succeeded
    assert normalized == [
        {
            "transaction_date": "2023-06-01",
            "country_code": "SL",
            "grant_code": "57801594",
            "budget_line_code": "4021014ON14A",
            "account_code": "52010",
            "site_code": "",
            "sector_code": "NUT",
            "transaction_code": "",
            "transaction_description": "",
            "currency_code": "USD",
            "budget_line_description": "DM - ON14A-Other nutrition",
            "amount": "39.88",
            "dummy_field_1": "Finance coordinator",
            "dummy_field_2": "6-CAM Support costs -2022",
            "dummy_field_3": "",
            "dummy_field_4": "6-CAM Support costs -2022",
            "dummy_field_5": "",
        }
    ]


def test_save_the_children_transaction_template_accepts_column_headers(monkeypatch):
    """A file that already carries column_N headers is used as-is."""
    _set_active_template(monkeypatch, "save_the_children")
    template = get_transaction_template()

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            ["SC-123", "BL-1", "Medical supplies", "5501", "202605", "1234.56", "", "CC-1"],
        ],
        analysis=_analysis_with_country("JO"),
    )

    assert succeeded
    assert normalized[0]["grant_code"] == "SC-123"
    assert normalized[0]["transaction_date"] == "2026-05-01"
    assert normalized[0]["dummy_field_1"] == "CC-1"
    # Columns absent from a short row simply stay empty.
    assert normalized[0]["dummy_field_5"] == ""


def test_save_the_children_transaction_template_takes_country_from_the_analysis(monkeypatch):
    """The layout has no country column, so the analysis supplies it."""
    _set_active_template(monkeypatch, "save_the_children")

    assert get_transaction_template().canonical_field_sources["country_code"] is None

    succeeded, normalized = normalize_uploaded_transaction_file(
        [["SC-123", "", "Medical supplies", "5501", "202605", "1234.56"]],
        analysis=_analysis_with_country("ET"),
    )

    assert succeeded
    assert normalized[0]["country_code"] == "ET"


def test_save_the_children_transaction_template_reports_missing_required_columns(monkeypatch):
    _set_active_template(monkeypatch, "save_the_children")
    template = get_transaction_template()

    succeeded, errors = normalize_uploaded_transaction_file(
        [
            ["column_1", "column_2"],
            ["SC-123", "BL-1"],
        ],
        analysis=_analysis_with_country(),
    )

    assert not succeeded
    assert errors == [
        "The Save the Children transaction file is missing required headers: "
        "column_3, column_4, column_5, column_6"
    ]
    assert set(template.required_source_fields) == {
        "column_1",
        "column_3",
        "column_4",
        "column_5",
        "column_6",
    }


def test_catholic_relief_services_transaction_template_loads(monkeypatch):
    _set_active_template(monkeypatch, "catholic_relief_services")
    template = get_transaction_template()

    assert template.id == "catholic_relief_services"
    assert template.label == "Catholic Relief Services"
    assert template.get_download_headers()[0] == "Date"
    assert "Amount" in template.get_download_headers()


def test_catholic_relief_services_transaction_template_normalizes_rows(monkeypatch):
    _set_active_template(monkeypatch, "catholic_relief_services")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "Date": "2025-11-30",
            "Country_code": "2624",
            "Grant_code": "30950",
            "Budget_line_code": "Allocated Direct Salaries - International",
            "Account_code": "601101",
            "Transaction_code": "26970701",
            "Transaction_description": "411619 135 601101 Allocations-20251209093818",
            "Currency_code": "USD",
            "Budget_line_description": "Allocated Direct Salaries - International",
            "Amount": "131.8",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized == [
        {
            "transaction_date": "2025-11-30",
            "country_code": "2624",
            "grant_code": "30950",
            "budget_line_code": "Allocated Direct Salaries - International",
            "account_code": "601101",
            "site_code": "",
            "sector_code": "",
            "transaction_code": "26970701",
            "transaction_description": "411619 135 601101 Allocations-20251209093818",
            "currency_code": "USD",
            "budget_line_description": "Allocated Direct Salaries - International",
            "amount": "131.8",
            "dummy_field_1": "",
            "dummy_field_2": "",
            "dummy_field_3": "",
            "dummy_field_4": "",
            "dummy_field_5": "",
        }
    ]


def test_catholic_relief_services_transaction_template_reports_missing_required_headers(monkeypatch):
    _set_active_template(monkeypatch, "catholic_relief_services")
    succeeded, errors = normalize_uploaded_transaction_file(
        [["Date", "Amount"]],
        analysis=object(),
    )

    assert not succeeded
    assert "Country_code" in errors[0]
    assert "Grant_code" in errors[0]


def test_accion_contra_el_hambre_transaction_template_loads(monkeypatch):
    _set_active_template(monkeypatch, "accion_contra_el_hambre")
    template = get_transaction_template()

    assert template.id == "accion_contra_el_hambre"
    assert template.label == "Accion Contra el Hambre"
    assert template.get_download_headers()[0] == "contract"
    assert "amount_eur" in template.get_download_headers()


@override_settings(ISO_CURRENCY_CODE="EUR")
def test_accion_contra_el_hambre_transaction_template_normalizes_rows(monkeypatch):
    _set_active_template(monkeypatch, "accion_contra_el_hambre")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "contract": "MLB2AS",
            "trans_date": "10/07/2024",
            "trans_no": "155885",
            "account": "600100",
            "cat5_adjusted": "BUD-200",
            "cat6": "2222",
            "description_ii": "Fuel and transport",
            "text": "2400LRECHARGE FUEL TOM CARTE GEN 43_TOTAL ENERGIES",
            "amount_eur": "2458.70",
            "dim_1": "ML",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized == [
        {
            "transaction_date": "2024-10-07",
            "country_code": "ML",
            "grant_code": "MLB2AS",
            "budget_line_code": "BUD-200",
            "account_code": "600100",
            "site_code": "",
            "sector_code": "2222",
            "transaction_code": "155885",
            "transaction_description": "2400LRECHARGE FUEL TOM CARTE GEN 43_TOTAL ENERGIES",
            "currency_code": "EUR",
            "budget_line_description": "Fuel and transport",
            "amount": "2458.70",
            "dummy_field_1": "",
            "dummy_field_2": "",
            "dummy_field_3": "",
            "dummy_field_4": "",
            "dummy_field_5": "",
        }
    ]


def test_accion_contra_el_hambre_transaction_template_reports_missing_required_headers(monkeypatch):
    _set_active_template(monkeypatch, "accion_contra_el_hambre")
    succeeded, errors = normalize_uploaded_transaction_file(
        [["contract", "amount_eur"]],
        analysis=object(),
    )

    assert not succeeded
    assert "trans_date" in errors[0]
    assert "account" in errors[0]
    assert "dim_1" in errors[0]


def test_cooperative_for_assistance_and_relief_everywhere_transaction_template_loads(monkeypatch):
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    template = get_transaction_template()

    assert template.id == "cooperative_for_assistance_and_relief_everywhere"
    assert template.label == "Cooperative for Assistance and Relief Everywhere (CARE)"
    assert template.get_download_headers()[0] == "journal_date"
    assert "bu_amount" in template.get_download_headers()


def test_cooperative_for_assistance_and_relief_everywhere_transaction_template_keeps_codes_as_text(
    monkeypatch,
):
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    template = get_transaction_template()

    assert template.get_download_field_type("budget_line_code") == TEXT_FIELD_TYPE
    assert template.get_download_field_type("account") == TEXT_FIELD_TYPE
    assert template.get_download_field_type("journal_id") == TEXT_FIELD_TYPE


def test_cooperative_for_assistance_and_relief_everywhere_transaction_template_normalizes_rows(
    monkeypatch,
):
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "journal_date": "03/31/2024",
            "country_code": "KEN01",
            "grant_code": "EF789",
            "budget_line_code": "5",
            "account": "501100",
            "site_code": "SI001",
            "project_id": "PROJAA0001",
            "journal_id": "0004044651",
            "transaction_description": "Indirect Cost Allocation",
            "currency_cd": "USD",
            "budget_line_description": "5 - LER program manager",
            "bu_amount": "103.14",
            "book_code": "COMN",
            "voucher_line_description": "Voucher description",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized == [
        {
            "transaction_date": "2024-03-31",
            "country_code": "KEN01",
            "grant_code": "EF789",
            "budget_line_code": "5",
            "account_code": "501100",
            "site_code": "SI001",
            "sector_code": "",
            "transaction_code": "0004044651",
            "transaction_description": "Indirect Cost Allocation",
            "currency_code": "USD",
            "budget_line_description": "5 - LER program manager",
            "amount": "103.14",
            "dummy_field_1": "",
            "dummy_field_2": "",
            "dummy_field_3": "",
            "dummy_field_4": "",
            "dummy_field_5": "",
        }
    ]


def test_cooperative_for_assistance_and_relief_everywhere_transaction_template_reports_missing_required_headers(
    monkeypatch,
):
    _set_active_template(monkeypatch, "cooperative_for_assistance_and_relief_everywhere")
    succeeded, errors = normalize_uploaded_transaction_file(
        [["journal_date", "bu_amount"], ["2024-03-31", "103.14"]],
        analysis=object(),
    )

    assert not succeeded
    assert "country_code" in errors[0]
    assert "grant_code" in errors[0]
    assert "account" in errors[0]


def test_danish_refugee_council_transaction_template_loads(monkeypatch):
    _set_active_template(monkeypatch, "danish_refugee_council")
    template = get_transaction_template()

    assert template.id == "danish_refugee_council"
    assert template.label == "Danish Refugee Council"
    assert template.get_download_headers()[0] == "column_1"
    assert template.get_download_headers()[-1] == "column_19"


def test_danish_refugee_council_transaction_template_normalizes_positional_rows(monkeypatch):
    _set_active_template(monkeypatch, "danish_refugee_council")
    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            [
                "30",
                "9",
                "2024",
                "0",
                "-",
                "2024-09-30",
                "2024-09-30",
                "BGD",
                "Gr00006009",
                "BGD-000312-06",
                "23011",
                "",
                "Shared Costs",
                "CA-BGD-000000514",
                "Salary Sept. 2024",
                "BDT",
                "BGD-082554 - 1.B.1 - Head of Programme (1)",
                "7823",
                "1",
            ]
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized == [
        {
            "transaction_date": "2024-09-30",
            "country_code": "BGD",
            "grant_code": "Gr00006009",
            "budget_line_code": "BGD-000312-06",
            "account_code": "23011",
            "site_code": "",
            "sector_code": "",
            "transaction_code": "",
            "transaction_description": "Salary Sept. 2024",
            "currency_code": "BDT",
            "budget_line_description": "BGD-082554 - 1.B.1 - Head of Programme (1)",
            "amount": "7823",
            "dummy_field_1": "",
            "dummy_field_2": "",
            "dummy_field_3": "",
            "dummy_field_4": "",
            "dummy_field_5": "",
        }
    ]


def test_danish_refugee_council_transaction_template_accepts_column_headers(monkeypatch):
    _set_active_template(monkeypatch, "danish_refugee_council")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "column_6": "2024-09-30",
            "column_8": "BGD",
            "column_9": "Gr00006009",
            "column_10": "BGD-000312-06",
            "column_11": "23011",
            "column_16": "BDT",
            "column_17": "BGD-082554 - 1.B.1 - Head of Programme (1)",
            "column_18": "7823",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=object(),
    )

    assert succeeded
    assert normalized[0]["transaction_date"] == "2024-09-30"
    assert normalized[0]["country_code"] == "BGD"
    assert normalized[0]["grant_code"] == "Gr00006009"
    assert normalized[0]["budget_line_code"] == "BGD-000312-06"
    assert normalized[0]["account_code"] == "23011"
    assert normalized[0]["currency_code"] == "BDT"
    assert normalized[0]["budget_line_description"] == "BGD-082554 - 1.B.1 - Head of Programme (1)"
    assert normalized[0]["amount"] == "7823"


def test_danish_refugee_council_transaction_template_reports_missing_required_column_headers(monkeypatch):
    _set_active_template(monkeypatch, "danish_refugee_council")
    succeeded, errors = normalize_uploaded_transaction_file(
        [["column_6", "column_18"], ["2024-09-30", "7823"]],
        analysis=object(),
    )

    assert not succeeded
    assert "column_8" in errors[0]
    assert "column_9" in errors[0]
    assert "column_11" in errors[0]


def test_mercy_corps_transaction_template_loads(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()

    assert template.id == "mercy_corps"
    assert template.label == "Mercy Corps"
    assert template.get_download_headers()[0] == "dwh_data_source"
    assert "usd_amount" in template.get_download_headers()


@override_settings(ISO_CURRENCY_CODE="USD")
def test_mercy_corps_transaction_template_normalizes_rows(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "posting_date": "2023-04-01",
            "posting_period": "2023-04",
            "calendar_quarter": "2023-2",
            "document_no": "XXXX",
            "g_l_account_no": "6910",
            "g_l_account_name": "Office Repairs & Maintenance",
            "fund_no": "999dummy",
            "dept_name": "ESTORIA DEPARTMENT",
            "office_code": "ES10",
            "lin_code": "91608L001",
            "lin_name": "Office Costs",
            "activity_code": "91608L001",
            "description": "FRAIS ENVOI OM DE 710 000XOF",
            "usd_amount": "9.85",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
        ],
        analysis=_analysis_with_country("SL"),
    )

    assert succeeded
    assert normalized == [
        {
            "transaction_date": "2023-04-01",
            "country_code": "SL",
            "grant_code": "999dummy",
            "budget_line_code": "91608L001",
            "account_code": "6910",
            "site_code": "ES10",
            "sector_code": "91608L001",
            "transaction_code": "XXXX",
            "transaction_description": "FRAIS ENVOI OM DE 710 000XOF",
            "currency_code": "USD",
            "budget_line_description": "Office Costs",
            "amount": "9.85",
            "dummy_field_1": "",
            "dummy_field_2": "",
            "dummy_field_3": "",
            "dummy_field_4": "",
            "dummy_field_5": "",
        }
    ]


def test_mercy_corps_transaction_template_uses_gl_account_name_when_lin_name_is_blank(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "posting_date": "2023-04-01",
            "g_l_account_no": "6910",
            "g_l_account_name": "Office Repairs & Maintenance",
            "fund_no": "999dummy",
            "lin_code": "91608L001",
            "lin_name": "",
            "usd_amount": "9.85",
        }
    )
    rows = [
        template.get_download_headers(),
        [source_row[header] for header in template.get_download_headers()],
    ]

    succeeded, normalized = normalize_uploaded_transaction_file(
        rows,
        analysis=_analysis_with_country("SL"),
    )
    errors = validate_uploaded_transaction_file(
        normalized,
        analysis=None,
        field_labels=template.get_validation_field_labels(rows),
        first_data_row=template.get_first_data_row_number(rows),
    )

    assert succeeded
    assert normalized[0]["country_code"] == "SL"
    assert normalized[0]["budget_line_description"] == "Office Repairs & Maintenance"
    assert errors == []


def test_mercy_corps_transaction_template_ignores_trailing_total_footer(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "posting_date": "2023-04-01",
            "document_no": "XXXX",
            "g_l_account_no": "6910",
            "fund_no": "999dummy",
            "dept_name": "ESTORIA DEPARTMENT",
            "office_code": "ES10",
            "lin_code": "91608L001",
            "lin_name": "Office Costs",
            "activity_code": "91608L001",
            "description": "FRAIS ENVOI OM DE 710 000XOF",
            "usd_amount": "9.85",
        }
    )

    succeeded, normalized = normalize_uploaded_transaction_file(
        [
            template.get_download_headers(),
            [source_row[header] for header in template.get_download_headers()],
            ["Total", *[""] * (len(template.get_download_headers()) - 1)],
            [""] * len(template.get_download_headers()),
            [],
        ],
        analysis=_analysis_with_country("SL"),
    )

    assert succeeded
    assert len(normalized) == 1
    assert normalized[0]["transaction_date"] == "2023-04-01"
    assert normalized[0]["amount"] == "9.85"


def test_mercy_corps_transaction_template_sets_icr_budget_line_description_before_validation(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()
    source_row = dict.fromkeys(template.get_download_headers(), "")
    source_row.update(
        {
            "posting_date": "2023-04-01",
            "g_l_account_no": "6910",
            "g_l_account_name": "Indirect Cost Recovery",
            "fund_no": "999dummy",
            "dept_name": "ESTORIA DEPARTMENT",
            "lin_code": "ICR",
            "lin_name": "",
            "usd_amount": "9.85",
        }
    )
    rows = [
        template.get_download_headers(),
        [source_row[header] for header in template.get_download_headers()],
    ]

    succeeded, normalized = normalize_uploaded_transaction_file(
        rows,
        analysis=_analysis_with_country("SL"),
    )
    errors = validate_uploaded_transaction_file(
        normalized,
        analysis=None,
        field_labels=template.get_validation_field_labels(rows),
        first_data_row=template.get_first_data_row_number(rows),
    )

    assert succeeded
    assert normalized[0]["budget_line_code"] == "ICR"
    assert normalized[0]["budget_line_description"] == "ICR"
    assert errors == []


def test_mercy_corps_transaction_template_reports_missing_required_headers(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    succeeded, errors = normalize_uploaded_transaction_file(
        [["posting_date", "usd_amount"], ["2023-04-01", "9.85"]],
        analysis=object(),
    )

    assert not succeeded
    assert "fund_no" in errors[0]
    assert "dept_name" not in errors[0]
    assert "g_l_account_no" in errors[0]
    assert "g_l_account_name" in errors[0]
    assert "lin_name" in errors[0]


def test_mercy_corps_validation_errors_use_source_field_and_column(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()
    rows = [
        [
            "fund_no",
            "g_l_account_no",
            "g_l_account_name",
            "lin_name",
            "usd_amount",
            "document_no",
            "description",
            "posting_date",
        ],
        [
            "999dummy",
            "6910",
            "Office Repairs & Maintenance",
            "Office Costs",
            "9.85",
            "XXXX",
            "FRAIS ENVOI OM DE 710 000XOF",
            "",
        ],
    ]

    succeeded, normalized = normalize_uploaded_transaction_file(
        rows,
        analysis=_analysis_with_country("SL"),
    )
    errors = validate_uploaded_transaction_file(
        normalized,
        analysis=None,
        field_labels=template.get_validation_field_labels(rows),
        first_data_row=template.get_first_data_row_number(rows),
    )

    assert succeeded
    assert errors == ["Row 2: posting_date (Column H) (Transaction Date) cannot be empty"]


def test_mercy_corps_invalid_character_errors_use_source_field_and_column(monkeypatch):
    _set_active_template(monkeypatch, "mercy_corps")
    template = get_transaction_template()
    uploaded_headers = [header.replace("_", " ").title() for header in template.get_download_headers()]
    source_row = dict.fromkeys(uploaded_headers, "")
    source_row.update(
        {
            "Posting Date": "2025-04-01",
            "Fund No": "999dummy",
            "Dept Name": "ESTORIA DEPARTMENT",
            "G L Account No": "5430",
            "Lin Name": "Conflicts Advisor",
            "Usd Amount": "1321.77",
            "Description": "486421-Alma Bezares Calder\ufffdn 03 31 25 FRG/32%",
        }
    )
    rows = [
        uploaded_headers,
        [source_row[header] for header in uploaded_headers],
    ]

    succeeded, normalized = normalize_uploaded_transaction_file(
        rows,
        analysis=_analysis_with_country("SL"),
    )
    errors = validate_uploaded_transaction_file(
        normalized,
        analysis=None,
        field_labels=template.get_validation_field_labels(rows),
        first_data_row=template.get_first_data_row_number(rows),
    )

    assert succeeded
    assert errors == [
        "Row 2: Description (Column AE) (Transaction Description) contains the Unicode replacement "
        "character (\ufffd, U+FFFD). This usually means the file was not decoded with the correct "
        "character encoding. Re-save or export the file as UTF-8, or correct the value in that cell."
    ]
